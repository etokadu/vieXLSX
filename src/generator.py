from __future__ import annotations

from dataclasses import asdict, dataclass
from datetime import datetime
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
import json
import os
import sys


TASK_HEADERS = [
    "ID",
    "Done",
    "Title",
    "Description",
    "Status",
    "Priority",
    "Due date",
    "Category",
    "Tags",
    "Created time",
    "Completed time"
]


@dataclass
class Task:
    id: int
    title: str
    done: bool = False
    description: str = ""
    status: str = "todo"
    priority: str = "normal"
    due_date: str = ""
    category: str = ""
    tags: str = ""
    created_time: str = ""
    completed_time: str = ""

    def toRow(self):
        return [
            self.id,
            "☑" if self.done else "☐",
            self.title,
            self.description,
            self.status,
            self.priority,
            self.due_date,
            self.category,
            self.tags,
            self.created_time,
            self.completed_time
        ]

    def toDict(self):
        return asdict(self)

    @classmethod
    def fromRow(cls, row, hasCheckbox=True):
        values = list(row) + [""] * len(TASK_HEADERS)

        if hasCheckbox:
            done = str(values[1] or "").lower() in (
                "☑", "true", "1", "yes", "x"
            )
            titleIndex = 2
        else:
            done = False
            titleIndex = 1

        return cls(
            id=int(values[0]),
            done=done,
            title=str(values[titleIndex] or ""),
            description=str(values[titleIndex + 1] or ""),
            status=str(values[titleIndex + 2] or "todo"),
            priority=str(values[titleIndex + 3] or "normal"),
            due_date=str(values[titleIndex + 4] or ""),
            category=str(values[titleIndex + 5] or ""),
            tags=str(values[titleIndex + 6] or ""),
            created_time=str(values[titleIndex + 7] or ""),
            completed_time=str(values[titleIndex + 8] or "")
        )


class ExcelGenerator:

    def __init__(self, path):
        self.path = path
        self.sheetName = "Tasks"
        self.filterSheetName = "site filter"
        self.tasks = []

        if os.path.isfile(self.path):
            self.wb = load_workbook(self.path)
        else:
            self.wb = Workbook()

        self.ws = self.getTaskSheet()
        self.syncFromExcel()

    def getTaskSheet(self):
        if self.sheetName in self.wb.sheetnames:
            return self.wb[self.sheetName]

        ws = self.wb.active
        ws.title = self.sheetName
        return ws

    def readExcelData(self):
        tasks = []
        firstRow = [
            cell.value
            for cell in self.ws[1]
        ]
        hasHeader = str(firstRow[0] or "").lower() == "id"
        hasCheckbox = any(
            str(cell or "").lower() == "done"
            for cell in firstRow
        )
        startRow = 2 if hasHeader else 1

        for row in self.ws.iter_rows(
            min_row=startRow,
            max_col=len(TASK_HEADERS),
            values_only=True
        ):
            if all(cell is None for cell in row):
                continue

            if row[0] is None:
                continue

            try:
                if hasHeader:
                    tasks.append(Task.fromRow(row, hasCheckbox))
                else:
                    tasks.append(Task(
                        id=int(row[0]),
                        title=str(row[1] or "")
                    ))
            except (TypeError, ValueError):
                continue

        return tasks

    def syncFromExcel(self):
        try:
            if os.path.isfile(self.path):
                self.wb = load_workbook(self.path)
                self.ws = self.getTaskSheet()

            self.tasks = self.readExcelData()
            return True
        except Exception:
            return False

    def nextId(self):
        return max(
            (task.id for task in self.tasks),
            default=0
        ) + 1

    def formatStyle(self, worksheet=None):
        try:
            worksheet = worksheet or self.ws
            headerFill = PatternFill(
                "solid",
                fgColor="1F4E78"
            )
            completedFill = PatternFill(
                "solid",
                fgColor="E2F0D9"
            )
            priorityFill = PatternFill(
                "solid",
                fgColor="FCE4D6"
            )
            mediumFill = PatternFill(
                "solid",
                fgColor="FFF2CC"
            )
            lowFill = PatternFill(
                "solid",
                fgColor="DDEBF7"
            )
            overdueFill = PatternFill(
                "solid",
                fgColor="FFC7CE"
            )
            widths = [8, 10, 28, 42, 14, 12, 16, 18, 24, 22, 22]

            checkboxValidation = DataValidation(
                type="list",
                formula1='"☐,☑"',
                allow_blank=False
            )
            checkboxValidation.error = "Choose ☐ or ☑."
            checkboxValidation.errorTitle = "Invalid checkbox"
            worksheet.add_data_validation(checkboxValidation)
            checkboxValidation.add(
                f"B2:B{max(2, worksheet.max_row)}"
            )

            appDirectory = (
                os.path.dirname(os.path.abspath(sys.executable))
                if getattr(sys, "frozen", False)
                else os.path.dirname(os.path.abspath(__file__))
            )
            configPath = os.path.join(appDirectory, "config.json")
            with open(configPath, "r", encoding="utf-8") as file:
                config = json.load(file)
            listsSheet = self.getListsSheet(config)
            categoryEnd = max(
                2,
                len(config.get("category_options", [])) + 1
            )
            tagEnd = max(
                2,
                len(config.get("tag_options", [])) + 1
            )
            categoryValidation = DataValidation(
                type="list",
                formula1=f"=Lists!$A$2:$A${categoryEnd}"
            )
            tagValidation = DataValidation(
                type="list",
                formula1=f"=Lists!$B$2:$B${tagEnd}"
            )
            worksheet.add_data_validation(categoryValidation)
            worksheet.add_data_validation(tagValidation)
            categoryValidation.add(
                f"H2:H{max(2, worksheet.max_row)}"
            )
            tagValidation.add(
                f"I2:I{max(2, worksheet.max_row)}"
            )

            for index, header in enumerate(TASK_HEADERS, start=1):
                cell = worksheet.cell(1, index, header)
                cell.fill = headerFill
                cell.font = Font(
                    color="FFFFFF",
                    bold=True
                )
                cell.alignment = Alignment(
                    horizontal="center",
                    vertical="center"
                )
                worksheet.column_dimensions[
                    cell.column_letter
                ].width = widths[index - 1]

            for row in worksheet.iter_rows(
                min_row=2,
                max_col=len(TASK_HEADERS)
            ):
                status = str(row[4].value or "").lower()
                priority = str(row[5].value or "").lower()
                dueDate = str(row[6].value or "")
                done = str(row[1].value or "") == "☑"

                for cell in row:
                    cell.alignment = Alignment(
                        vertical="top",
                        wrap_text=True
                    )

                if status in ("done", "completed", "complete"):
                    for cell in row:
                        cell.fill = completedFill
                        cell.font = Font(
                            color="666666",
                            strike=True
                        )
                elif priority in ("high", "urgent", "critical"):
                    for cell in row:
                        cell.fill = priorityFill
                elif priority in ("medium", "normal"):
                    for cell in row:
                        cell.fill = mediumFill
                elif priority == "low":
                    for cell in row:
                        cell.fill = lowFill

                if done:
                    row[1].font = Font(
                        color="008000",
                        bold=True
                    )

                try:
                    dueDateValue = datetime.fromisoformat(
                        dueDate
                    ).date()
                    if not done and dueDateValue < datetime.now().date():
                        for cell in row:
                            cell.fill = overdueFill
                except ValueError:
                    pass

            worksheet.freeze_panes = "A2"
            worksheet.auto_filter.ref = (
                f"A1:K{max(1, worksheet.max_row)}"
            )
            return True
        except Exception:
            return False

    def getListsSheet(self, config):
        if "Lists" in self.wb.sheetnames:
            worksheet = self.wb["Lists"]
            worksheet.delete_rows(1, worksheet.max_row)
        else:
            worksheet = self.wb.create_sheet("Lists")

        worksheet.append(["Categories", "Tags"])
        categories = config.get("category_options", [])
        tags = config.get("tag_options", [])
        for index in range(max(len(categories), len(tags))):
            worksheet.append([
                categories[index] if index < len(categories) else "",
                tags[index] if index < len(tags) else ""
            ])

        worksheet.sheet_state = "hidden"
        return worksheet

    def saveTasks(self, tasks=None, sheetName=None):
        try:
            tasks = self.tasks if tasks is None else tasks
            sheetName = self.sheetName if sheetName is None else sheetName
            parentPath = os.path.dirname(os.path.abspath(self.path))

            if parentPath:
                os.makedirs(parentPath, exist_ok=True)

            if sheetName in self.wb.sheetnames:
                worksheet = self.wb[sheetName]
                worksheet.delete_rows(1, worksheet.max_row)
            else:
                worksheet = self.wb.create_sheet(sheetName)

            worksheet.append(TASK_HEADERS)

            for task in tasks:
                worksheet.append(task.toRow())

            if not self.formatStyle(worksheet):
                return False

            if sheetName == self.sheetName:
                self.ws = worksheet
                self.saveSettingsSheet()
            self.wb.save(self.path)
            return True
        except Exception:
            return False

    def updateExcel(self):
        return self.saveTasks(self.tasks, self.sheetName)

    def saveSettingsSheet(self):
        try:
            configPath = os.path.join(
                os.path.dirname(os.path.abspath(__file__)),
                "config.json"
            )
            with open(configPath, "r", encoding="utf-8") as file:
                config = json.load(file)

            settingsName = "Settings"
            if settingsName in self.wb.sheetnames:
                worksheet = self.wb[settingsName]
                worksheet.delete_rows(1, worksheet.max_row)
            else:
                worksheet = self.wb.create_sheet(settingsName)

            worksheet.append(["Setting", "Value"])
            worksheet.append(["App", config.get("app_name", "vieXLSX")])
            worksheet.append(["Version", config.get("version", "1.0.0")])
            worksheet.append(["Author", config.get("author", "Kadu")])
            worksheet.append(["GitHub", config.get("github", "")])
            worksheet.append(["Last update", config.get("last_update", "")])
            worksheet.append([
                "New latest feature",
                config.get("latest_feature", "")
            ])
            worksheet.column_dimensions["A"].width = 24
            worksheet.column_dimensions["B"].width = 80
            self.formatSettingsStyle(worksheet)
            taskIndex = self.wb.sheetnames.index(self.sheetName)
            settingsIndex = self.wb.sheetnames.index(settingsName)
            if settingsIndex != taskIndex + 1:
                self.wb._sheets.insert(
                    taskIndex + 1,
                    self.wb._sheets.pop(settingsIndex)
                )
            return True
        except Exception:
            return False

    def formatSettingsStyle(self, worksheet):
        headerFill = PatternFill(
            "solid",
            fgColor="1F4E78"
        )

        for cell in worksheet[1]:
            cell.fill = headerFill
            cell.font = Font(
                color="FFFFFF",
                bold=True
            )
            cell.alignment = Alignment(
                horizontal="center"
            )

        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )

    def sortTasks(self, field, reverse=False):
        try:
            if field not in Task.__dataclass_fields__:
                return False

            self.tasks.sort(
                key=lambda task: str(
                    getattr(task, field)
                ).lower(),
                reverse=reverse
            )
            return self.updateExcel()
        except Exception:
            return False

    def filterTasks(self, criteria):
        try:
            result = []

            for task in self.tasks:
                matched = all(
                    str(value).lower()
                    in str(getattr(task, key, "")).lower()
                    if key == "tags"
                    else str(getattr(task, key, "")).lower()
                    == str(value).lower()
                    for key, value in criteria.items()
                    if value not in (None, "")
                )

                if matched:
                    result.append(task)

            if not self.saveTasks(
                result,
                self.filterSheetName
            ):
                return None

            return result
        except Exception:
            return None

    def searchTasks(self, query):
        query = str(query).lower()
        return [
            task
            for task in self.tasks
            if query in " ".join(
                str(value)
                for value in task.toDict().values()
            ).lower()
        ]

    def importData(self, data, formatRequired=False):
        try:
            lines = [
                line.strip()
                for line in data.split("\n")
                if line.strip()
            ]

            for line in lines:
                parts = line.split("|", 1)
                self.tasks.append(Task(
                    id=int(parts[0])
                    if formatRequired
                    else self.nextId(),
                    title=parts[1]
                    if formatRequired and len(parts) > 1
                    else line
                ))

            return True
        except Exception:
            return False
